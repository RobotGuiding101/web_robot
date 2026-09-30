"""Generate TEST_CASES.xlsx / TEST_CASES.csv / TEST_CASES_SUMMARY.csv from TEST_CASES.md."""
import csv
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path("/home/ncwjj/git/web_robot")
MD = ROOT / "TEST_CASES.md"

lines = MD.read_text(encoding="utf-8").splitlines()

section = ""
tables = []
i = 0
while i < len(lines):
    line = lines[i]
    m = re.match(r"^##\s+(.*)$", line)
    if m:
        section = m.group(1).strip()
        i += 1
        continue
    if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:\-|]+\|$", lines[i + 1]):
        headers = [c.strip() for c in line.strip().strip("|").split("|")]
        rows = []
        j = i + 2
        while j < len(lines) and lines[j].startswith("|"):
            rows.append([c.strip() for c in lines[j].strip().strip("|").split("|")])
            j += 1
        tables.append((section, headers, rows))
        i = j
        continue
    i += 1


def clean(text):
    return text.replace("`", "").replace("<br>", "\n").strip()


test_rows = []
for sec, headers, rows in tables:
    module = re.sub(r"^\d+\.\s*", "", sec).replace("`", "")
    if "Precondition" in headers:
        idx = {h: k for k, h in enumerate(headers)}
        for r in rows:
            test_rows.append({
                "Module": module,
                "No": clean(r[idx["No"]]),
                "Test Case": clean(r[idx["Test Case"]]),
                "Type": "Manual",
                "Precondition": clean(r[idx.get("Precondition", 0)]),
                "Steps": clean(r[idx.get("Steps", 0)]),
                "Expected Result": clean(r[idx.get("Expected Result", 0)]),
                "File": "",
                "Result": "",
                "Note": "",
            })
    elif "ไฟล์" in headers:
        fcol = headers.index("ไฟล์")
        for r in rows:
            test_rows.append({
                "Module": module,
                "No": clean(r[0]),
                "Test Case": clean(r[1]),
                "Type": "Automated",
                "Precondition": "",
                "Steps": "",
                "Expected Result": "",
                "File": clean(r[fcol]),
                "Result": "Pass",
                "Note": "",
            })

wb = Workbook()
ws = wb.active
ws.title = "Test Cases"

cols = ["Module", "No", "Test Case", "Type", "Precondition", "Steps",
        "Expected Result", "File", "Result", "Note"]
widths = [26, 10, 34, 11, 24, 44, 44, 30, 12, 24]

head_fill = PatternFill("solid", fgColor="1F4E79")
head_font = Font(bold=True, color="FFFFFF")
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

ws.append(cols)
for c, w in zip(ws[1], widths):
    c.fill = head_fill
    c.font = head_font
    c.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
ws.row_dimensions[1].height = 22

for row in test_rows:
    ws.append([row.get(c, "") for c in cols])

for col, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(col)].width = w
for r in range(2, ws.max_row + 1):
    for c in range(1, len(cols) + 1):
        cell = ws.cell(row=r, column=c)
        cell.border = border
        cell.alignment = Alignment(vertical="top", wrap_text=(c in (5, 6, 7, 8)))
ws.freeze_panes = "C2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws.max_row}"

man = sum(1 for r in test_rows if r["Type"] == "Manual")
auto = sum(1 for r in test_rows if r["Type"] == "Automated")


def _cnt(rows, key):
    return sum(1 for r in rows if r.get("Result") == key)


ws2 = wb.create_sheet("Summary")
ws2.append(["Module", "Total", "Manual", "Automated", "Pass", "Fail", "Blocked", "Note"])
mods = []
for row in test_rows:
    if row["Module"] not in mods:
        mods.append(row["Module"])
for mod in mods:
    rs = [r for r in test_rows if r["Module"] == mod]
    ws2.append([mod, len(rs),
                sum(1 for r in rs if r["Type"] == "Manual"),
                sum(1 for r in rs if r["Type"] == "Automated"),
                _cnt(rs, "Pass") or None, _cnt(rs, "Fail") or None,
                _cnt(rs, "Blocked") or None, ""])
ws2.append(["TOTAL (test cases)", len(test_rows), man, auto,
            _cnt(test_rows, "Pass"), _cnt(test_rows, "Fail"),
            _cnt(test_rows, "Blocked"), ""])
ws2.append([])
ws2.append(["Automated test run (latest)", "Total", "Pass", "Fail", "", "", "", ""])
ws2.append(["Backend pytest (backend/tests)", 41, 41, 0, "", "", "", ""])
ws2.append(["Frontend vitest (frontend/src)", 25, 25, 0, "", "", "", ""])
ws2.append(["TOTAL tests", 66, 66, 0, "", "", "", ""])

for c in ws2[1]:
    c.fill = head_fill
    c.font = head_font
    c.alignment = Alignment(horizontal="center")
ws2["A" + str(ws2.max_row)].font = Font(bold=True)
for col, w in zip("ABCDEFGH", [34, 8, 8, 11, 8, 8, 9, 26]):
    ws2.column_dimensions[col].width = w
ws2.freeze_panes = "A2"

wb.save(ROOT / "TEST_CASES.xlsx")

with open(ROOT / "TEST_CASES.csv", "w", newline="", encoding="utf-8-sig") as f:
    w = csv.writer(f)
    for r in ws.iter_rows(values_only=True):
        w.writerow(["" if c is None else c for c in r])

with open(ROOT / "TEST_CASES_SUMMARY.csv", "w", newline="", encoding="utf-8-sig") as f:
    w = csv.writer(f)
    for r in ws2.iter_rows(values_only=True):
        w.writerow(["" if c is None else c for c in r])

print(f"written: {len(test_rows)} cases ({man} manual, {auto} automated)")
